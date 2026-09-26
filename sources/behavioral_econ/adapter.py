"""Classic judgment and decision problems with published human answer distributions, from public
replication data: Many Labs 1, 2 and 3 (OSF, CC0 / CC-BY) and the 19-country prospect theory replication
(Ruggeri et al. 2020, OSF esxc4). Every question carries the replication sample's item-level choice shares."""

from __future__ import annotations

import csv
import io
import zipfile
from collections import Counter
from pathlib import Path
from typing import Iterator

import httpx
import openpyxl
import polars as pl

from askjev.model import HumanDist, Question

NAME = "behavioral_econ"
MIN_COUNTRY_N = 100  # per-country populations need at least this many answers

FILES = {
    "ml1/Datasets.zip": "https://osf.io/download/nqg97/",  # Many Labs 1 (OSF wx7ck, CC0)
    "ml2/ML2_S1.csv": "https://osf.io/download/cwjp3/",  # Many Labs 2 slate 1 (OSF 8cd4r)
    "ml2/ML2_S2.csv": "https://osf.io/download/jg9hc/",  # Many Labs 2 slate 2
    "ml2/ML2_SourceInfo.xlsx": "https://osf.io/download/uv4qx/",  # site -> country
    "ml3/ML3_PPool_MTurk.csv": "https://osf.io/download/tdgk9/",  # Many Labs 3 (OSF ct89g)
    "pt19/pt_data.csv": "https://osf.io/download/p2f7w/",  # prospect theory replication (OSF esxc4)
}
ML1_TSV = "Data/Tab.delimited.Cleaned.dataset.WITH.variable.labels.dat"

LIC = {
    "ml1": "CC0 1.0 (Many Labs 1 data, Klein et al. 2014, OSF wx7ck)",
    "ml2": "CC0 1.0 (Many Labs 2 data, Klein et al. 2018, OSF 8cd4r)",
    "ml3": "CC-BY 4.0 (Many Labs 3 data, Ebersole et al. 2016, OSF ct89g)",
    "pt19": "No license set on OSF; public data, research use with citation (Ruggeri et al. 2020, Nature Human Behaviour 4:622-633, OSF esxc4)",
}
POP = {
    "ml1": "Many Labs 1 participants (36 samples, 12 countries)",
    "ml2": "Many Labs 2 participants (125 samples, 36 countries)",
    "ml3": "Many Labs 3 participants (20 US university pools + MTurk)",
    "pt19": "Prospect theory replication participants (19 countries)",
}
CITE = {
    "ml1": "Many Labs 1 (Klein et al. 2014, Social Psychology 45:142-152), item-level responses",
    "ml2": "Many Labs 2 (Klein et al. 2018, AMPPS 1:443-490), item-level responses",
    "ml3": "Many Labs 3 (Ebersole et al. 2016, JESP 67:68-82), item-level responses",
    "pt19": "Ruggeri et al. 2020, Replicating patterns of prospect theory for decision under risk, item-level responses",
}

RISK = "self.personality.risk_decision_style"
SACRIFICE = "self.values.sacrificial_dilemmas"


def fetch(raw_dir: Path) -> None:
    for rel, url in FILES.items():
        out = raw_dir / rel
        if out.exists():
            continue
        out.parent.mkdir(parents=True, exist_ok=True)
        r = httpx.get(url, follow_redirects=True, timeout=600)
        r.raise_for_status()
        out.write_bytes(r.content)
    tsv = raw_dir / "ml1" / ML1_TSV
    if not tsv.exists():
        with zipfile.ZipFile(raw_dir / "ml1" / "Datasets.zip") as zf:
            zf.extract(ML1_TSV, raw_dir / "ml1")
            zf.extract("Data/Data released under CC0 license.txt", raw_dir / "ml1")


# ---------------------------------------------------------------------------------------------------------
# Item specs. `col` is the dataset column(s) (pooled when several); `map` sends a raw answer to an option
# key (choice), "true"/"false" (noul) or a level index (score, several raw values may share a level when a
# long rating scale is binned to at most 7 concrete levels).
# ---------------------------------------------------------------------------------------------------------

AGREE7 = ["1", "2", "3", "4", "5", "6", "7"]


def _levels(n: int, bins: list[list[int]] | None = None) -> dict[str, int]:
    bins = bins or [[i] for i in range(1, n + 1)]
    return {str(v): i for i, b in enumerate(bins) for v in b}


BIN9 = [[1, 2], [3, 4], [5], [6, 7], [8, 9]]
BIN10 = [[1, 2], [3, 4], [5, 6], [7, 8], [9, 10]]
BIN11 = [[1, 2], [3, 4], [5, 6, 7], [8, 9], [10, 11]]

DISEASE = (
    "Imagine that your country is preparing for the outbreak of an unusual disease, which is expected to kill "
    "600 people. Two alternative programs to combat the disease have been proposed, and the exact scientific "
    "estimates of their consequences are as follows. "
)
SUNK = (
    "Imagine that your favorite football team is playing an important game. You have a ticket to the game that "
    "you {how}. However, on the day of the game, it happens to be freezing cold. What do you do?"
)
SUNK_LEVELS = [
    "I stay at home and skip the game without a second thought",
    "I lean toward staying at home, though I think about going",
    "I am torn and could just as well stay home or go",
    "I lean toward bundling up and going to the game",
    "I go to the game without a second thought, cold or not",
]
TV_LOW = {
    "Up to a half hour": "up_to_half_an_hour",
    "Half an hour to an hour": "half_an_hour_to_an_hour",
    "One to one and a half hours": "one_to_one_and_a_half_hours",
    "One and a half to two hours": "one_and_a_half_to_two_hours",
    "Two to two and a half hours": "two_to_two_and_a_half_hours",
    "More than two and a half hours": "more_than_two_and_a_half_hours",
}
TV_HIGH = {
    "Up to two and a half hours": "up_to_two_and_a_half_hours",
    "Two and a half to three hours": "two_and_a_half_to_three_hours",
    "Three to three and a half hours": "three_to_three_and_a_half_hours",
    "Three and a half to four hours": "three_and_a_half_to_four_hours",
    "Four to four and a half hours": "four_to_four_and_a_half_hours",
    "More than four and a half hours": "more_than_four_and_a_half_hours",
}
YESNO_ML1 = {"Yes": "true", "No": "false"}
YESNO_12 = {"1": "true", "2": "false"}

ROSS1 = (
    "As you are leaving your neighborhood supermarket, a man in a business suit asks you whether you like shopping "
    "in that store. You reply quite honestly that you do like shopping there and that, besides being close to your "
    "home, the supermarket seems to have very good meats and produce at reasonably low prices. The man then reveals "
    "that a videotape crew has filmed your comments and asks you to sign a release allowing them to use the unedited "
    "film for a TV commercial that the supermarket chain is preparing. Would you sign the release or refuse to sign it?"
)
ROSS2 = (
    "While driving through a rural area near your home, you are stopped by a county police officer who informs you "
    "that you have been clocked by radar at 38 miles per hour in a 25-mph zone. You believe this information to be "
    "accurate. After the officer leaves, you inspect your citation and find that the details on the summons about "
    "weather, visibility, time, and location of the violation are highly inaccurate. The citation says you may "
    "either pay an $80 fine by mail without appearing in court, or appear in municipal court within the next two "
    "weeks to contest the charge. Would you pay the $80 fine by mail or appear in court?"
)
HAUS = {
    "denise": (
        "Denise is a passenger on a train whose driver has just shouted that the train's brakes have failed, and who "
        "then fainted from the shock. On the track ahead are five people; the banks are so steep that they will not "
        "be able to get off the track in time. The track has a side track leading off to the right, and Denise can "
        "turn the train onto it. Unfortunately there is one person on the right-hand track. Denise can turn the "
        "train, killing the one; or she can refrain from turning the train, letting the five die. Is it morally "
        "permissible for Denise to switch the train to the side track?"
    ),
    "frank": (
        "Frank is on a footbridge over the train tracks. He knows trains and can see that the one approaching the "
        "bridge is out of control. On the track under the bridge there are five people; the banks are so steep that "
        "they will not be able to get off the track in time. Frank knows that the only way to stop an out-of-control "
        "train is to drop a very heavy weight into its path. But the only available, sufficiently heavy weight is a "
        "large man wearing a backpack, also watching the train from the footbridge. Frank can shove the man with the "
        "backpack onto the track in the path of the train, killing him; or he can refrain from doing this, letting "
        "the five die. Is it morally permissible for Frank to shove the man?"
    ),
    "ned": (
        "Ned is taking his daily walk near the train tracks when he notices that the approaching train is out of "
        "control: the driver saw five men walking across the tracks and slammed on the brakes, but the brakes failed "
        "and the men will not be able to get off the tracks in time. Ned is standing next to a switch that will "
        "temporarily turn the train onto a side track. There is a heavy object on the side track: if the train hits "
        "it, the object will slow the train down and give the men time to escape. Unfortunately, the heavy object is "
        "a man, standing on the side track with his back turned. Ned can throw the switch, preventing the train from "
        "killing the five men but killing the man, or he can refrain, letting the five die. Is it morally "
        "permissible for Ned to throw the switch?"
    ),
    "oscar": (
        "Oscar is taking his daily walk near the train tracks when he notices that the approaching train is out of "
        "control: the driver saw five men walking across the tracks and slammed on the brakes, but the brakes failed "
        "and the driver fainted, and the men will not be able to get off the track in time. Oscar is standing next "
        "to a switch that will temporarily turn the train onto a side track. There is a heavy object on the side "
        "track: if the train hits it, the object will slow the train down and give the men time to escape. "
        "Unfortunately, there is a man standing on the side track in front of the heavy object, with his back "
        "turned. Oscar can throw the switch, preventing the train from killing the five men but killing the man, or "
        "he can refrain, letting the five die. Is it morally permissible for Oscar to throw the switch?"
    ),
}
PERMISSIBLE = {"true": "Yes, it is morally permissible", "false": "No, it is not morally permissible"}
TVER = (
    "Imagine that you are about to buy a ceramic vase for ${vase} and a wall hanging for ${wall}. The salesperson "
    "tells you that the wall hanging you want is on sale for ${sale} at the store's other branch, a 20-minute drive "
    "away. Would you make the trip to the other branch?"
)
SHAFIR = (
    "Imagine that you serve on the jury of an only-child sole-custody case following a relatively messy divorce. "
    "The facts are complicated by ambiguous economic, social, and emotional considerations, and you decide to base "
    "your decision entirely on these observations. Parent A: average income, average health, average working hours, "
    "reasonable rapport with the child, relatively stable social life. Parent B: above-average income, very close "
    "relationship with the child, extremely active social life, lots of work-related travel, minor health problems. "
)
PARENTS = {
    "parent_a": "Average income, average health, average working hours, reasonable rapport with the child, relatively stable social life",
    "parent_b": "Above-average income, very close relationship with the child, extremely active social life, lots of work-related travel, minor health problems",
}
DIRECTOR = (
    "A director was working on a music video. His assistant said: 'I took a look at the first cut of your video, and "
    "it looks to me like some of the images in it will encourage {who} to French kiss in public.' The director said: "
    "'Look, I know that it will be encouraging {who} to French kiss in public, but I don't care at all about that. I "
    "just want to make a video that will increase sales of the album.' He included the images in the video. Sure "
    "enough, it encouraged {who} to French kiss in public. "
)
INTENT_LEVELS = [
    "He did not do it intentionally at all",
    "He almost certainly did not do it intentionally",
    "He probably did not do it intentionally",
    "It is a toss-up whether he did it intentionally",
    "He probably did it intentionally",
    "He almost certainly did it intentionally",
    "He definitely did it intentionally",
]
WRONG_VIDEO_LEVELS = [
    "Making the video was not wrong at all",
    "Making the video was barely wrong, a slight lapse at most",
    "Making the video was a little wrong",
    "Making the video was somewhat wrong",
    "Making the video was fairly wrong",
    "Making the video was clearly wrong",
    "Making the video was definitely wrong",
]
CHAIRMAN = (
    "The vice-president of a company went to the chairman of the board and said, 'We are thinking of starting a new "
    "program. It will help us increase profits, {and_} it will also {effect} the environment.' The chairman of the "
    "board answered, 'I don't care at all about {effecting} the environment. I just want to make as much profit as I "
    "can. Let's start the new program.' They started the new program. Sure enough, the environment was {effected}. "
)
AGREE_LEVELS = [
    "I strongly disagree with the statement",
    "I moderately disagree with the statement",
    "I slightly disagree with the statement",
    "I neither agree nor disagree with the statement",
    "I slightly agree with the statement",
    "I moderately agree with the statement",
    "I strongly agree with the statement",
]
PRAISE_LEVELS = [
    "He deserves no praise at all",
    "He deserves a token word of praise",
    "He deserves a little praise",
    "He deserves a fair amount of praise",
    "He deserves real praise",
    "He deserves a great deal of praise",
    "He deserves a lot of praise, like a hero of the story",
]
BLAME_LEVELS = [
    "He deserves no blame at all",
    "He deserves a token bit of blame",
    "He deserves a little blame",
    "He deserves a fair amount of blame",
    "He deserves real blame",
    "He deserves a great deal of blame",
    "He deserves a lot of blame, like the villain of the story",
]
GIFT = (
    "Imagine you are about to leave the country and have received a goodbye gift from a friend. It is a wool {item}, "
    "from a nearby department store. The store carries a variety of wool {item}s. The worst costs ${lo} and the best "
    "costs ${hi}. The one your friend bought you costs ${paid}. How generous do you think your friend was?"
)
GENEROUS_LEVELS = [
    "My friend was not generous at all",
    "My friend was barely generous",
    "My friend was a little generous",
    "My friend was moderately generous",
    "My friend was quite generous",
    "My friend was very generous",
    "My friend was extremely generous",
]
LECTURE = (
    "Imagine that you are in a large lecture with a few hundred students, sitting in the middle section a little "
    "more than half-way back. The professor asks a question about the readings, but no one raises a hand to answer. "
    "{prep} The class sits in silence for two minutes before the professor explains that if no one volunteers, he "
    "will choose someone. How likely is it that the professor will call on you?"
)
CALL_LEVELS = [
    "It is very unlikely that he will call on me",
    "It is somewhat unlikely that he will call on me",
    "It is about as likely as not that he will call on me",
    "It is somewhat likely that he will call on me",
    "It is very likely that he will call on me",
]
LIFE_SAT_LEVELS = [
    "I am deeply dissatisfied with my life",
    "I am more dissatisfied than satisfied with my life",
    "I am neither especially satisfied nor dissatisfied with my life",
    "I am fairly satisfied with my life",
    "I am thoroughly satisfied with my life",
]
PARTNER_LEVELS = [
    "I am deeply dissatisfied with my relationship",
    "I am more dissatisfied than satisfied with my relationship",
    "I am neither especially satisfied nor dissatisfied with my relationship",
    "I am fairly satisfied with my relationship",
    "I am thoroughly satisfied with my relationship",
]
GW_BELIEF = [
    "I am not at all convinced that global warming is happening",
    "I am somewhat doubtful that global warming is happening",
    "I am fairly convinced that global warming is happening",
    "I am completely convinced that global warming is happening",
]
GW_WORRY = [
    "I do not worry about global warming at all",
    "I worry about global warming a little",
    "I worry about global warming quite a bit",
    "I worry about global warming a great deal",
]
DISGUST3 = [
    "It would not bother me at all",
    "It would gross me out a little",
    "It would disgust me a lot",
]
TIPI_LEVELS = [
    "I strongly disagree that this describes me",
    "I moderately disagree that this describes me",
    "I slightly disagree that this describes me",
    "I neither agree nor disagree that this describes me",
    "I slightly agree that this describes me",
    "I moderately agree that this describes me",
    "I strongly agree that this describes me",
]
TIPI = [
    ("Extraverted, enthusiastic", "self.personality.big_five.extraversion.sociability_energy"),
    ("Critical, quarrelsome", "self.personality.big_five.agreeableness.trust_modesty_temper"),
    ("Dependable, self-disciplined", "self.personality.big_five.conscientiousness.drive_self_discipline"),
    ("Anxious, easily upset", "self.personality.big_five.neuroticism.fear_worry"),
    ("Open to new experiences, complex", "self.personality.big_five.openness.ideas_imagination"),
    ("Reserved, quiet", "self.personality.big_five.extraversion.sociability_energy"),
    ("Sympathetic, warm", "self.personality.big_five.agreeableness.kindness_cooperation"),
    ("Disorganized, careless", "self.personality.big_five.conscientiousness.order_caution"),
    ("Calm, emotionally stable", "self.personality.big_five.neuroticism.temper_moodiness"),
    ("Conventional, uncreative", "self.personality.big_five.openness.ideas_imagination"),
]
SISE_LEVELS = [
    "This is not true of me at all",
    "This is mostly not true of me",
    "This is partly true of me",
    "This is mostly true of me",
    "This is very true of me",
]
TRUEFALSE = {"1": "true", "2": "false"}


def _items() -> list[dict]:
    it: list[dict] = []

    def add(**kw):
        it.append(kw)

    # ---- Many Labs 1 -----------------------------------------------------------------------------------------
    add(sid="ml1:diseaseframinga", data="ml1", col=["diseaseframinga"], primitive="choice", node=RISK, kind="personality",
        text=DISEASE + "If Program A is adopted, 200 people will be saved. If Program B is adopted, there is a 1/3 "
        "probability that 600 people will be saved and a 2/3 probability that no people will be saved. Which "
        "program would you choose?",
        options={"program_a": "200 people will be saved", "program_b": "A 1/3 probability that 600 people will be saved and a 2/3 probability that no one will be saved"},
        map={"200 people will be saved": "program_a", "1/3 probability to save all, 2/3 nobody will be saved": "program_b"},
        meta={"study": "Tversky & Kahneman 1981 Asian disease problem, gain frame"})
    add(sid="ml1:diseaseframingb", data="ml1", col=["diseaseframingb"], primitive="choice", node=RISK, kind="personality",
        text=DISEASE + "If Program A is adopted, 400 people will die. If Program B is adopted, there is a 1/3 "
        "probability that nobody will die and a 2/3 probability that 600 people will die. Which program would you choose?",
        options={"program_a": "400 people will die", "program_b": "A 1/3 probability that nobody will die and a 2/3 probability that 600 people will die"},
        map={"400 people will die": "program_a", "1/3 probability nobody will die, 2/3 that 600 will die": "program_b"},
        meta={"study": "Tversky & Kahneman 1981 Asian disease problem, loss frame"})
    for v, how in (("a", "have paid handsomely for"), ("b", "received for free from a friend")):
        add(sid=f"ml1:sunkcost{v}", data="ml1", col=[f"sunkcost{v}"], primitive="score", node=RISK, kind="personality",
            text=SUNK.format(how=how), options=SUNK_LEVELS, map=_levels(9, BIN9),
            meta={"study": "Oppenheimer et al. 2009 sunk cost", "scale": "1 (stay at home) to 9 (go to game), binned 1-2/3-4/5/6-7/8-9"})
    add(sid="ml1:allowedforbiddena", data="ml1", col=["allowedforbiddena"], primitive="noul", node="self.values.fairness_justice",
        kind="values", text="Do you think the United States should forbid public speeches against democracy?",
        options={"true": "Yes, it should forbid them", "false": "No"}, map=YESNO_ML1, flags=["political"],
        meta={"study": "Rugg 1941 allowed/forbidden asymmetry, forbid wording"})
    add(sid="ml1:allowedforbiddenb", data="ml1", col=["allowedforbiddenb"], primitive="noul", node="self.values.fairness_justice",
        kind="values", text="Do you think the United States should allow public speeches against democracy?",
        options={"true": "Yes, it should allow them", "false": "No"}, map=YESNO_ML1, flags=["political"],
        meta={"study": "Rugg 1941 allowed/forbidden asymmetry, allow wording"})
    add(sid="ml1:reciprocityusa", data="ml1", col=["reciprocityusa"], primitive="noul", node="self.values.fairness_justice",
        kind="values", text="Do you think the United States should let newspaper reporters from North Korea come in "
        "and send back to their papers the news as they see it?",
        options={"true": "Yes", "false": "No"}, map=YESNO_ML1, flags=["political"],
        meta={"study": "Hyman & Sheatsley 1950 norm of reciprocity, asked first"})
    add(sid="ml1:reciprocityotherb", data="ml1", col=["reciprocityotherb"], primitive="noul", node="self.values.fairness_justice",
        kind="values", text="Do you think North Korea should let American newspaper reporters come in and send back "
        "to America the news as they see it?",
        options={"true": "Yes", "false": "No"}, map=YESNO_ML1, flags=["political"],
        meta={"study": "Hyman & Sheatsley 1950 norm of reciprocity, asked first"})
    for v, m in (("a", TV_LOW), ("b", TV_HIGH)):
        add(sid=f"ml1:scales{v}", data="ml1", col=[f"scales{v}"], primitive="choice", node="self.lifestyle.screens_media",
            kind="personality", text="How much television do you watch per day?",
            options={k: lab for lab, k in m.items()}, map=m,
            meta={"study": f"Schwarz et al. 1985 low vs high category scales, {'low' if v == 'a' else 'high'}-range options"})

    # ---- Many Labs 2 -----------------------------------------------------------------------------------------
    add(sid="ml2:rott1.1", data="ml2s1", col=["rott1.1"], primitive="choice", node=RISK, kind="personality",
        text="Imagine you have the opportunity either to meet and kiss your favorite movie star or to receive $50 in "
        "cash. Which would you prefer?",
        options={"meet_and_kiss_my_favorite_movie_star": None, "receive_50_dollars_in_cash": None},
        map={"1": "meet_and_kiss_my_favorite_movie_star", "2": "receive_50_dollars_in_cash"},
        meta={"study": "Rottenstreich & Hsee 2001, certainty condition"})
    add(sid="ml2:rott2.1", data="ml2s1", col=["rott2.1"], primitive="choice", node=RISK, kind="personality",
        text="Imagine you can take part either in a lottery that offers a 1% chance to meet and kiss your favorite "
        "movie star or in a lottery that offers a 1% chance to receive $50 in cash. Which lottery would you prefer?",
        options={"1_percent_chance_to_kiss_my_favorite_movie_star": None, "1_percent_chance_of_50_dollars_cash": None},
        map={"1": "1_percent_chance_to_kiss_my_favorite_movie_star", "2": "1_percent_chance_of_50_dollars_cash"},
        meta={"study": "Rottenstreich & Hsee 2001, 1% chance condition"})
    add(sid="ml2:ross.s1.2", data="ml2s1", col=["ross.s1.2"], primitive="choice", node=RISK, kind="personality",
        text=ROSS1, options={"sign_the_release": None, "refuse_to_sign": None},
        map={"1": "sign_the_release", "2": "refuse_to_sign"}, meta={"study": "Ross, Greene & House 1977 false consensus, supermarket"})
    add(sid="ml2:ross.s2.2", data="ml2s2", col=["ross.s2.2"], primitive="choice", node=RISK, kind="personality",
        text=ROSS2, options={"pay_the_fine_by_mail": None, "appear_in_court": None},
        map={"1": "pay_the_fine_by_mail", "2": "appear_in_court"}, meta={"study": "Ross, Greene & House 1977 false consensus, speeding ticket"})
    for sid, data, key in (("ml2:haus1.1", "ml2s1", "denise"), ("ml2:haus2.1", "ml2s1", "frank"),
                           ("ml2:haus3.1", "ml2s2", "ned"), ("ml2:haus4.1", "ml2s2", "oscar")):
        add(sid=sid, data=data, col=[sid.split(":")[1]], primitive="noul", node=SACRIFICE, kind="values",
            text=HAUS[key], options=PERMISSIBLE, map=YESNO_12, meta={"study": f"Hauser et al. 2007 trolley problems ({key})"})
    for sid, vase, wall, sale in (("ml2:tver1.1", 30, 250, 240), ("ml2:tver2.1", 250, 30, 20)):
        add(sid=sid, data="ml2s2", col=[sid.split(":")[1]], primitive="noul", node="self.lifestyle.money_habits",
            kind="personality", text=TVER.format(vase=vase, wall=wall, sale=sale),
            options={"true": "Yes, I would drive to the other branch", "false": "No, I would buy it here"}, map=YESNO_12,
            meta={"study": "Tversky & Kahneman 1981 mental accounting ($10 off a {}-dollar item)".format(wall)})
    add(sid="ml2:shaf1.1", data="ml2s2", col=["shaf1.1"], primitive="choice", node="self.values.fairness_justice", kind="values",
        text=SHAFIR + "To which parent would you award sole custody of the child?", options=PARENTS,
        map={"1": "parent_a", "2": "parent_b"}, meta={"study": "Shafir 1993 choosing vs rejecting, award"})
    add(sid="ml2:shaf2.1", data="ml2s2", col=["shaf2.1"], primitive="choice", node="self.values.fairness_justice", kind="values",
        text=SHAFIR + "Which parent would you deny sole custody of the child?", options=PARENTS,
        map={"1": "parent_a", "2": "parent_b"}, meta={"study": "Shafir 1993 choosing vs rejecting, deny"})
    for n, who, flags in (("1", "homosexual men", ["political"]), ("2", "couples", [])):
        pre = DIRECTOR.format(who=who)
        add(sid=f"ml2:inba{n}.3", data="ml2s1", col=[f"inba{n}.3"], primitive="score", node="self.values.everyday_ethics",
            kind="values", text=pre + f"Did the director intentionally encourage {who} to French kiss in public?",
            options=INTENT_LEVELS, map=_levels(7), flags=flags, meta={"study": "Inbar et al. 2009 intentionality (Knobe effect)"})
        add(sid=f"ml2:inba{n}.4", data="ml2s1", col=[f"inba{n}.4"], primitive="noul", node="self.love.etiquette_social_norms",
            kind="values", text=f"Is there anything wrong with {who} French kissing in public?",
            options={"true": "Yes, something is wrong with it", "false": "No, nothing is wrong with it"}, map=YESNO_12,
            flags=flags, meta={"study": "Inbar et al. 2009"})
        add(sid=f"ml2:inba{n}.5", data="ml2s1", col=[f"inba{n}.5"], primitive="score", node="self.values.everyday_ethics",
            kind="values", text=pre + f"Was it wrong of the director to make a video that he knew would encourage {who} to "
            "French kiss in public?", options=WRONG_VIDEO_LEVELS, map=_levels(7), flags=flags,
            meta={"study": "Inbar et al. 2009", "scale": "1 (not at all) to 7 (definitely)"})
    help_ = CHAIRMAN.format(and_="and", effect="help", effecting="helping", effected="helped")
    harm = CHAIRMAN.format(and_="but", effect="harm", effecting="harming", effected="harmed")
    add(sid="ml2:knob1.3", data="ml2s2", col=["knob1.3"], primitive="score", node="self.values.everyday_ethics", kind="values",
        text=help_ + 'How much do you agree with the statement "The chairman helped the environment intentionally"?',
        options=AGREE_LEVELS, map=_levels(7), meta={"study": "Knobe 2003 side-effect effect, help"})
    add(sid="ml2:knob2.3", data="ml2s2", col=["knob2.3"], primitive="score", node="self.values.everyday_ethics", kind="values",
        text=harm + 'How much do you agree with the statement "The chairman harmed the environment intentionally"?',
        options=AGREE_LEVELS, map=_levels(7), meta={"study": "Knobe 2003 side-effect effect, harm"})
    add(sid="ml2:knob1.4", data="ml2s2", col=["knob1.4"], primitive="score", node="self.values.everyday_ethics", kind="values",
        text=help_ + "How much praise does the chairman deserve for what he did?", options=PRAISE_LEVELS, map=_levels(7),
        meta={"study": "Knobe 2003 side-effect effect, praise"})
    add(sid="ml2:knob2.4", data="ml2s2", col=["knob2.4"], primitive="score", node="self.values.everyday_ethics", kind="values",
        text=harm + "How much blame does the chairman deserve for what he did?", options=BLAME_LEVELS, map=_levels(7),
        meta={"study": "Knobe 2003 side-effect effect, blame"})
    for sid, item, lo, hi, paid in (("ml2:hsee1.1", "scarf", 10, 100, 90), ("ml2:hsee2.1", "coat", 100, "1,000", 110)):
        add(sid=sid, data="ml2s2", col=[sid.split(":")[1]], primitive="score", node="self.love.friendship", kind="evaluative",
            text=GIFT.format(item=item, lo=lo, hi=hi, paid=paid), options=GENEROUS_LEVELS,
            map={str(i): i for i in range(7)}, meta={"study": "Hsee 1998 less is better", "scale": "0-6"})
    for sid, prep in (("ml2:rise1.3", "You have not done the reading and feel confident that you would not be able to answer the question."),
                      ("ml2:rise2.3", "You have done the reading and feel confident that the professor would like your answer, but you prefer not to volunteer answers in large classes.")):
        add(sid=sid, data="ml2s2", col=[sid.split(":")[1]], primitive="score", node="self.mind.luck_fate.superstitions_omens",
            kind="forecast", text=LECTURE.format(prep=prep), options=CALL_LEVELS, map=_levels(10, BIN10),
            meta={"study": "Risen & Gilovich 2008 tempting fate", "scale": "1-10, binned in pairs"})
    add(sid="ml2:schw1.1", data="ml2s2", col=["schw1.1"], primitive="score", node="self.mind.happiness_wellbeing", kind="personality",
        text="How satisfied are you currently with your life as a whole?", options=LIFE_SAT_LEVELS, map=_levels(11, BIN11),
        meta={"study": "Schwarz, Strack & Mai 1991, asked first", "scale": "1-11, binned 1-2/3-4/5-7/8-9/10-11"})
    add(sid="ml2:schw2.1", data="ml2s2", col=["schw2.1"], primitive="score", node="self.love.romance_partnership", kind="personality",
        text="Think about your relationship with your partner (spouse or date). How satisfied are you currently with your relationship?",
        options=PARTNER_LEVELS, map=_levels(11, BIN11),
        meta={"study": "Schwarz, Strack & Mai 1991, asked first", "scale": "1-11, binned 1-2/3-4/5-7/8-9/10-11"})
    add(sid="ml2:zav.dv.2", data="ml2s2", col=["zav.dv.2"], primitive="score", node="self.values.animals_environment", kind="values",
        text="How convinced are you that global warming is happening?", options=GW_BELIEF, map=_levels(4), flags=["political"],
        meta={"study": "Zaval et al. 2014 (primes pooled)"})
    add(sid="ml2:zav.dv.3", data="ml2s2", col=["zav.dv.3"], primitive="score", node="self.values.animals_environment", kind="values",
        text="How much do you personally worry about global warming?", options=GW_WORRY, map=_levels(4), flags=["political"],
        meta={"study": "Zaval et al. 2014 (primes pooled)"})
    for c, stmt in (("disg1.11", "I never let any part of my body touch the toilet seat in a public washroom."),
                    ("disg1.12", "I probably would not go to my favorite restaurant if I found out that the cook had a cold.")):
        add(sid=f"ml2:{c}", data="ml2s1", col=[c], primitive="noul", node="self.personality.emotions_stress", kind="personality",
            text=f'Is this statement true of you: "{stmt}"', options={"true": "True of me", "false": "False of me"},
            map=TRUEFALSE, meta={"study": "Disgust Scale-Revised, contamination (Olatunji et al. 2007)"})
    for c, exp, flags in (("disg2.10", "You take a sip of soda and realize that you drank from the glass that an acquaintance of yours had been drinking from.", []),
                          ("disg2.12", "A friend offers you a piece of chocolate shaped like dog poop.", []),
                          ("disg2.13", "As part of a sex education class, you are required to inflate a new lubricated condom using your mouth.", ["sensitive"])):
        add(sid=f"ml2:{c}", data="ml2s1", col=[c], primitive="score", node="self.personality.emotions_stress", kind="personality",
            text=f'How disgusting would you find this experience: "{exp}"', options=DISGUST3, map=_levels(3), flags=flags,
            meta={"study": "Disgust Scale-Revised (Olatunji et al. 2007)"})
    for i, (trait, node) in enumerate(TIPI, 1):
        add(sid=f"ml2+ml3:tipi_{i}", data="ml2both+ml3", col=[f"tipi_{i}"], ml3col=f"big5_{i:02d}", primitive="score", node=node,
            kind="personality", text=f'How well does this describe you: "I see myself as {trait.lower()}"?',
            options=TIPI_LEVELS, map=_levels(7), meta={"study": "Ten-Item Personality Inventory (Gosling et al. 2003)"})
    add(sid="ml2:sise", data="ml2both", col=["sise"], primitive="score", node="self.personality.self_concept", kind="personality",
        text='How true of you is this statement: "I have high self-esteem"?', options=SISE_LEVELS, map=_levels(5),
        meta={"study": "Single-Item Self-Esteem scale (Robins et al. 2001)"})
    add(sid="ml2:subjwell", data="ml2both", col=["subjwell"], primitive="score", node="self.mind.happiness_wellbeing", kind="personality",
        text="All things considered, how satisfied are you with your life as a whole these days?", options=LIFE_SAT_LEVELS,
        map=_levels(10, BIN10), meta={"study": "Veenhoven 2009 life satisfaction", "scale": "1-10, binned in pairs"})

    # ---- Many Labs 3 -----------------------------------------------------------------------------------------
    for letter in "KLNRV":
        add(sid=f"ml3:{letter.lower()}position", data="ml3", col=[f"{letter.lower()}position"], primitive="choice",
            node="self.mind.epistemics", kind="perception",
            text=f"Think of English words that contain the letter {letter}. Does {letter} appear more often as the first "
            "letter of a word or as the third letter?",
            options={"as_the_first_letter": None, "as_the_third_letter": None},
            map={"1": "as_the_first_letter", "2": "as_the_third_letter"},
            meta={"study": "Tversky & Kahneman 1973 availability heuristic (letter position)",
                  "note": "no truth set: the classic claim (third position more frequent) depends on the word counts used"})

    # ---- Prospect theory replication (US-version dollar amounts) --------------------------------------------
    PT = [
        ("1", "Which would you prefer: a 33% chance of gaining $5,000, a 66% chance of gaining $4,800 and a 1% chance of $0; or $4,800 for sure?",
         ("the_gamble", "A 33% chance of gaining $5,000, a 66% chance of gaining $4,800, and a 1% chance of $0"), ("4800_dollars_for_sure", "Guaranteed $4,800")),
        ("2", "Which would you prefer: a 33% chance of gaining $5,000 (67% chance of $0), or a 34% chance of gaining $4,800 (66% chance of $0)?",
         ("33_percent_chance_of_5000", "A 33% chance of gaining $5,000 (67% chance of $0)"), ("34_percent_chance_of_4800", "A 34% chance of gaining $4,800 (66% chance of $0)")),
        ("3", "Which would you prefer: an 80% chance of gaining $8,000 (20% chance of $0), or $6,000 for sure?",
         ("80_percent_chance_of_8000", "An 80% chance of gaining $8,000 (20% chance of $0)"), ("6000_dollars_for_sure", "A 100% guarantee of gaining $6,000")),
        ("4", "Which would you prefer: a 20% chance of gaining $8,000 (80% chance of $0), or a 25% chance of gaining $6,000 (75% chance of $0)?",
         ("20_percent_chance_of_8000", "A 20% chance of gaining $8,000 (80% chance of $0)"), ("25_percent_chance_of_6000", "A 25% chance of gaining $6,000 (75% chance of $0)")),
        ("5", "Which would you prefer: a 45% chance of gaining $12,000 (55% chance of $0), or a 90% chance of gaining $6,000 (10% chance of $0)?",
         ("45_percent_chance_of_12000", "A 45% chance of gaining $12,000 (55% chance of $0)"), ("90_percent_chance_of_6000", "A 90% chance of gaining $6,000 (10% chance of $0)")),
        ("6", "Which would you prefer: a 0.1% chance of gaining $12,000 (99.9% chance of $0), or a 0.2% chance of gaining $6,000 (99.8% chance of $0)?",
         ("0_1_percent_chance_of_12000", "A 0.1% chance of gaining $12,000 (99.9% chance of $0)"), ("0_2_percent_chance_of_6000", "A 0.2% chance of gaining $6,000 (99.8% chance of $0)")),
        ("7", "Which would you prefer: an 80% chance of losing $8,000 (20% chance of losing nothing), or losing $6,000 for sure?",
         ("80_percent_chance_to_lose_8000", "An 80% chance of losing $8,000 (20% chance of losing $0)"), ("lose_6000_dollars_for_sure", "A 100% guarantee of losing $6,000")),
        ("8", "Which would you prefer: a 20% chance of losing $8,000 (80% chance of losing nothing), or a 25% chance of losing $6,000 (75% chance of losing nothing)?",
         ("20_percent_chance_to_lose_8000", "A 20% chance of losing $8,000 (80% chance of losing $0)"), ("25_percent_chance_to_lose_6000", "A 25% chance of losing $6,000 (75% chance of losing $0)")),
        ("9", "Which would you prefer: a 45% chance of losing $12,000 (55% chance of losing nothing), or a 90% chance of losing $6,000 (10% chance of losing nothing)?",
         ("45_percent_chance_to_lose_12000", "A 45% chance of losing $12,000 (55% chance of losing $0)"), ("90_percent_chance_to_lose_6000", "A 90% chance of losing $6,000 (10% chance of losing $0)")),
        ("10", "Which would you prefer: a 0.1% chance of losing $12,000 (99.9% chance of losing nothing), or a 0.2% chance of losing $6,000 (99.8% chance of losing nothing)?",
         ("0_1_percent_chance_to_lose_12000", "A 0.1% chance of losing $12,000 (99.9% chance of losing $0)"), ("0_2_percent_chance_to_lose_6000", "A 0.2% chance of losing $6,000 (99.8% chance of losing $0)")),
        ("11", "Imagine you are playing a game with two levels, but you have to make a choice about the second level before you know the outcome "
         "of the first. At the first level, there is a 75% chance that the game will end without you winning anything, and a 25% chance that "
         "you will advance to the second level. What would you choose for the second level: an 80% chance of gaining $8,000 (20% chance of $0), or $6,000 for sure?",
         ("80_percent_chance_of_8000", "An 80% chance of gaining $8,000 (20% chance of $0)"), ("6000_dollars_for_sure", "A 100% guarantee of gaining $6,000")),
        ("12", "Imagine you were given $2,000 right now to play a game. Which would you prefer: a 50% chance to gain an additional $2,000 (50% chance of gaining nothing more), or an additional $1,000 for sure?",
         ("50_percent_chance_of_2000_more", "A 50% chance to gain an additional $2,000 (50% chance of gaining $0 beyond what you already have)"), ("1000_more_for_sure", "A 100% guarantee of gaining an additional $1,000")),
        ("13", "Imagine you were given $4,000 right now to play a game. Which would you prefer: a 50% chance of losing $2,000 (50% chance of losing nothing), or losing $1,000 for sure?",
         ("50_percent_chance_to_lose_2000", "A 50% chance you will lose $2,000 (50% chance of losing $0)"), ("lose_1000_dollars_for_sure", "A 100% chance you will lose $1,000")),
        ("14", "Which would you prefer: a 25% chance of gaining $12,000 (75% chance of $0), or a 25% chance of gaining $8,000 plus a 25% chance of gaining $4,000 (50% chance of $0)?",
         ("25_percent_chance_of_12000", "A 25% chance of gaining $12,000 (75% chance of $0)"), ("25_percent_8000_or_25_percent_4000", "A 25% chance of gaining $8,000, a 25% chance of gaining $4,000, and a 50% chance of $0")),
        ("15", "Which would you prefer: a 25% chance of losing $12,000 (75% chance of losing nothing), or a 25% chance of losing $8,000 plus a 25% chance of losing $4,000 (50% chance of losing nothing)?",
         ("25_percent_chance_to_lose_12000", "A 25% chance of losing $12,000 (75% chance of losing $0)"), ("25_percent_lose_8000_or_25_percent_lose_4000", "A 25% chance of losing $8,000, a 25% chance of losing $4,000, and a 50% chance of losing $0")),
        ("16", "Which would you prefer: a 0.1% chance of gaining $10,000 (99.9% chance of $0), or $10 for sure?",
         ("0_1_percent_chance_of_10000", "A 0.1% chance of gaining $10,000 (99.9% chance of $0)"), ("10_dollars_for_sure", "A 100% guarantee of gaining $10")),
        ("17", "Which would you prefer: a 0.1% chance of losing $10,000 (99.9% chance of losing nothing), or losing $10 for sure?",
         ("0_1_percent_chance_to_lose_10000", "A 0.1% chance of losing $10,000 (99.9% chance of losing $0)"), ("lose_10_dollars_for_sure", "A 100% guarantee of losing $10")),
    ]
    for col, text, (ka, da), (kb, db) in PT:
        add(sid=f"pt19:item{col}", data="pt19", col=[col], primitive="choice", node=RISK, kind="personality", text=text,
            options={ka: da, kb: db}, map={"1.0": ka, "0.0": kb},
            meta={"study": "Kahneman & Tversky 1979 prospect theory problem, 2020 19-country replication",
                  "note": "amounts are the US survey's; other countries saw the same structure in local currency"})
    return it


# ---------------------------------------------------------------------------------------------------------
# Data loading: each dataset -> list of (row dict of str, country label or None).
# ---------------------------------------------------------------------------------------------------------

ML1_COUNTRY = {"US": "United States", "PL": "Poland", "CA": "Canada", "IT": "Italy", "MY": "Malaysia", "BR": "Brazil",
               "TR": "Turkey", "CN": "China", "IN": "India", "NL": "Netherlands", "CZ": "Czech Republic", "UK": "United Kingdom"}


def _load(raw_dir: Path) -> dict[str, list[tuple[dict, str | None]]]:
    out: dict[str, list] = {}
    ml1 = pl.read_csv(raw_dir / "ml1" / ML1_TSV, separator="\t", infer_schema_length=0, quote_char=None, encoding="utf8-lossy")
    out["ml1"] = [(r, ML1_COUNTRY.get((r.get("citizenship") or "").strip())) for r in ml1.iter_rows(named=True)]

    ws = openpyxl.load_workbook(raw_dir / "ml2" / "ML2_SourceInfo.xlsx", read_only=True).active
    rows = list(ws.iter_rows(values_only=True))
    hdr = rows[0]
    si, ci = hdr.index("Source"), hdr.index("Country")
    country = {r[si]: r[ci] for r in rows[1:] if r[si]}
    for slate in ("s1", "s2"):
        df = pl.read_csv(raw_dir / "ml2" / f"ML2_{slate.upper()}.csv", infer_schema_length=0)
        out[f"ml2{slate}"] = [(r, country.get(r["source"])) for r in df.iter_rows(named=True)]
    out["ml2both"] = out["ml2s1"] + out["ml2s2"]

    t = (raw_dir / "ml3" / "ML3_PPool_MTurk.csv").read_text(encoding="latin-1").replace("\r\n", "\n").replace("\r", "\n")
    out["ml3"] = [(r, None) for r in csv.DictReader(io.StringIO(t))]

    pt = pl.read_csv(raw_dir / "pt19" / "pt_data.csv", infer_schema_length=0)
    out["pt19"] = [(r, r["Country"]) for r in pt.iter_rows(named=True)]
    return out


def _dist(counts: Counter, keys: list[str]) -> dict[str, float]:
    n = sum(counts.values())
    return {k: counts.get(k, 0) / n for k in keys}


def _human(spec: dict, rows: list[tuple[dict, str | None]], cols: list[str], study: str, keys: list[str]) -> list[HumanDist]:
    total: Counter = Counter()
    by_country: dict[str, Counter] = {}
    for r, ctry in rows:
        for c in cols:
            v = (r.get(c) or "").strip()
            if v in spec["map"]:
                k = str(spec["map"][v])
                total[k] += 1
                if ctry:
                    by_country.setdefault(ctry, Counter())[k] += 1
    if not total:
        return []
    hs = [HumanDist(population=POP[study], distribution=_dist(total, keys), n=sum(total.values()), source=CITE[study])]
    for ctry in sorted(by_country):
        cnt = by_country[ctry]
        if sum(cnt.values()) >= MIN_COUNTRY_N and len(by_country) > 1:
            hs.append(HumanDist(population=f"{POP[study]}: {ctry}", distribution=_dist(cnt, keys), n=sum(cnt.values()), source=CITE[study]))
    return hs


def normalize(raw_dir: Path) -> Iterator[Question]:
    data = _load(raw_dir)
    for spec in _items():
        p = spec["primitive"]
        if p == "choice":
            keys = list(spec["options"])
        elif p == "noul":
            keys = ["true", "false"]
        else:
            keys = [str(i) for i in range(len(spec["options"]))]
        human: list[HumanDist] = []
        for part in spec["data"].split("+"):
            if part == "ml3":
                human += _human(spec, data["ml3"], [spec.get("ml3col") or spec["col"][0]], "ml3", keys)
            else:
                study = "ml2" if part.startswith("ml2") else part
                human += _human(spec, data[part], spec["col"], study, keys)
        if not human:
            raise ValueError(f"no responses for {spec['sid']}")
        study0 = spec["data"].split("+")[0]
        lic = "; ".join(LIC["ml2" if s.startswith("ml2") else s] for s in spec["data"].split("+"))
        flags = spec.get("flags") or []
        yield Question(
            text=spec["text"],
            primitive=p,
            hemisphere="self",
            kind=spec["kind"],
            origin="dataset",
            source=NAME,
            options=spec["options"],
            node_hint=spec["node"],
            source_item_id=spec["sid"],
            license=lic,
            human=human,
            meta={"dataset": study0, "columns": spec["col"]} | spec.get("meta", {}) | ({"flags": flags} if flags else {}),
        )
