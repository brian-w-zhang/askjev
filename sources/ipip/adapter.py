"""IPIP Big-Five items (public domain) from ipip.ori.org scoring keys, with Open Psychometrics
response distributions for the 50 Big-Five Factor Markers."""

from __future__ import annotations

import html
import io
import random
import re
import zipfile
from pathlib import Path
from typing import Iterator

import httpx
import polars as pl

from askjev.model import HumanDist, Question
from askjev.sampling import env_int, hash_order

NAME = "ipip"
LICENSE = "Public domain (IPIP items)"
NORMS_LICENSE = "Open Psychometrics raw data (no license stated)"
TARGET = 600  # the original sample (kept byte-identical)
TARGET_ALL = env_int("TARGET_IPIP", TARGET)  # expansion: the rest of the 3,320-item pool, appended after
EXTRA_SALT = "ipip-extra-v1"
SEED = 20260924

BASE = "https://ipip.ori.org/"
PAGES = {
    "markers": "newBigFive5broadKey.htm",  # 50- and 100-item Big-Five Factor Markers
    "neo": "newNEOKey.htm",  # IPIP-NEO 30 facets (300 items)
    "bfas": "BFASKeys.htm",  # Big Five Aspect Scales (100 items)
    "ab5c": "newAB5CKey.htm",  # AB5C 45 facets
    "hexaco": "newHEXACO_PI_key.htm",  # IPIP HEXACO-PI analog (last-resort fill)
}
# Expansion sources: the full alphabetical item list (3,320 items with survey codes) and Tedone's item
# assignment table (item -> instrument, construct label, key) linked from ipip.ori.org/ItemAssignmentTable.htm.
ALPHA_PAGE = "AlphabeticalItemList.htm"
TEDONE_XLSX = "TedoneItemAssignmentTable30APR21.xlsx"
NORMS_URL = "https://openpsychometrics.org/_rawdata/IPIP-FFM-data-8Nov2018.zip"
NORMS_DIR = "IPIP-FFM-data-8Nov2018"

LEVELS = [
    "This does not describe me at all",
    "This describes me a little",
    "This describes me moderately well",
    "This describes me well",
    "This describes me very well",
]
TRAITS = ["openness", "conscientiousness", "extraversion", "agreeableness", "neuroticism"]
ROMAN = {"I": "extraversion", "II": "agreeableness", "III": "conscientiousness", "IV": "neuroticism", "V": "openness"}
NEO_DOMAIN = {"N": "neuroticism", "E": "extraversion", "O": "openness", "A": "agreeableness", "C": "conscientiousness"}
BFAS_DOMAIN = {
    "Neuroticism": "neuroticism",
    "Agreeableness": "agreeableness",
    "Conscientiousness": "conscientiousness",
    "Extraversion": "extraversion",
    "Openness/Intellect": "openness",
}
# Items whose stem is not a verb phrase, so "I " + stem would not be a sentence.
STEM_FIX = {"willing to": "am willing to", "never at a loss": "am never at a loss", "interested in": "am interested in"}
FLIP = {"+": "-", "-": "+"}
HEXACO_TRAIT = {"X": "extraversion", "A": "agreeableness", "C": "conscientiousness", "O": "openness"}


def fetch(raw_dir: Path) -> None:
    headers = {"User-Agent": "Mozilla/5.0 (askjev research)"}
    for page in [*PAGES.values(), ALPHA_PAGE, TEDONE_XLSX]:
        out = raw_dir / page
        if not out.exists():
            r = httpx.get(BASE + page, headers=headers, follow_redirects=True, timeout=60)
            r.raise_for_status()
            out.write_bytes(r.content)
    if not (raw_dir / NORMS_DIR / "data-final.csv").exists():
        r = httpx.get(NORMS_URL, headers=headers, follow_redirects=True, timeout=600)
        r.raise_for_status()
        (raw_dir / Path(NORMS_URL).name).write_bytes(r.content)
        zipfile.ZipFile(io.BytesIO(r.content)).extractall(raw_dir)


# ---------- parsing ----------

def _lines(path: Path) -> list[str]:
    t = path.read_bytes().decode("cp1252", errors="replace")
    t = re.sub(r"\s+", " ", t)
    t = re.sub(r"(?is)<(script|style).*?</\1>", " ", t)
    t = re.sub(r"(?i)</?(p|td|tr|br|h\d|li|div|table)\b[^>]*>", "\n", t)
    t = html.unescape(re.sub(r"<[^>]+>", " ", t))
    return [ln for ln in (re.sub(r"\s+", " ", x).strip() for x in t.split("\n")) if ln]


def _sign(line: str) -> str | None:
    if re.fullmatch(r"\+ ?keyed", line):
        return "+"
    if re.fullmatch(r"[–—-] ?keyed", line):
        return "-"
    return None


def _is_item(line: str) -> bool:
    return bool(re.fullmatch(r"\[?[A-Z][A-Za-z'’ ,\-\"“”/]*[.!?][\"”]?\]?( [a-d])?", line)) and not line.startswith("Note")


def _parse_markers(path: Path):
    trait = size = sign = None
    for ln in _lines(path):
        m = re.match(r"Factor ([IV]+) ", ln)
        if m:
            trait = ROMAN[m.group(1)]
            continue
        m = re.match(r"(10|20)-item scale", ln)
        if m:
            size = 50 if m.group(1) == "10" else 100
            continue
        if (s := _sign(ln)) is not None:
            sign = s
            continue
        if trait and sign and _is_item(ln):
            # Factor IV is keyed for Emotional Stability; flip to neuroticism.
            k = FLIP[sign] if trait == "neuroticism" else sign
            scale = f"Big-Five Factor Markers ({size}-item)"
            yield ln, trait, k, None, scale


def _parse_neo(path: Path):
    trait = facet = sign = None
    for ln in _lines(path):
        if ln.startswith("5 NEO Domains"):
            break
        m = re.match(r"([NEOAC])(\d): ([A-Z\- ]+?) \(", ln)
        if m:
            trait, facet, sign = NEO_DOMAIN[m.group(1)], f"{m.group(1)}{m.group(2)} {m.group(3).title()}", None
            continue
        if (s := _sign(ln)) is not None:
            sign = s
            continue
        if trait and sign and _is_item(ln):
            yield ln, trait, sign, facet, "IPIP-NEO-300"


def _parse_bfas(path: Path):
    trait = aspect = sign = None
    for ln in _lines(path):
        if ln.startswith("Note."):
            break
        if ln in BFAS_DOMAIN:
            trait, aspect, sign = BFAS_DOMAIN[ln], None, None
            continue
        if re.fullmatch(r"[A-Z][a-z]+", ln) and trait:
            aspect, sign = ln, None
            continue
        if (s := _sign(ln)) is not None:
            sign = s
            continue
        if trait and aspect and sign and _is_item(ln):
            yield ln, trait, sign, aspect, "BFAS"


def _parse_ab5c(path: Path):
    trait = facet = sign = None
    for ln in _lines(path):
        if ln.startswith("Note."):
            break
        m = re.match(r"([IV]+)\+/([IV]+)([+-]) vs .*?\(.*?\)\s*([A-Z][A-Z\- ]+)$", ln)
        if m:
            trait = ROMAN[m.group(1)]
            facet = f"{m.group(1)}+/{m.group(2)}{m.group(3)} {m.group(4).strip().title()}"
            sign = None
            continue
        if (s := _sign(ln)) is not None:
            sign = s
            continue
        # Bracketed items belong to other facets; skip them here.
        if trait and sign and _is_item(ln) and not ln.startswith("["):
            k = FLIP[sign] if trait == "neuroticism" else sign
            yield ln, trait, k, facet, "AB5C"


def _parse_hexaco(path: Path):
    """IPIP HEXACO-PI analog. X/A/C/O map to the Big Five trait of the same name; H and E have no
    Big Five counterpart and get trait None (placed under the Big Five parent)."""
    dom = facet = sign = None
    for ln in _lines(path):
        if ln.startswith("Note"):
            break
        m = re.match(r"(.+?) \(([HEXACO]):\w+\)", ln)
        if m:
            dom, facet, sign = m.group(2), m.group(1).strip(), None
            continue
        if (s := _sign(ln)) is not None:
            sign = s
            continue
        if dom and sign and _is_item(ln):
            yield ln, HEXACO_TRAIT.get(dom), sign, f"{dom}:{facet}", "IPIP-HEXACO"


def _sentence(stem: str) -> str:
    s = re.sub(r" [a-d]$", "", stem.strip()).strip("[]").strip()
    s = s.replace("’", "'").replace("“", "'").replace("”", "'").replace('"', "'")
    s = re.sub(r"\s+", " ", s)
    low = s[:1].lower() + s[1:]
    for k, v in STEM_FIX.items():
        if low.startswith(k):
            low = v + low[len(k):]
    s = "I " + low
    if not s.endswith((".", "!", "?", ".'")):
        s += "."
    return s


def _norm(s: str) -> str:
    return re.sub(r"[^a-z ]", "", s.lower().replace("'", "")).strip()


# ---------- norms ----------

def _norms(raw_dir: Path) -> dict[str, tuple[dict[str, float], int, str]]:
    """Per-item response shares (1-5 -> "0".."4") over respondents with IPC==1 and all 50 answers valid."""
    d = raw_dir / NORMS_DIR
    items = {}
    for ln in (d / "codebook.txt").read_text(encoding="utf-8", errors="replace").splitlines():
        m = re.match(r"^((?:EXT|EST|AGR|CSN|OPN)\d+)\t(.+)$", ln)
        if m:
            items[m.group(1)] = m.group(2).strip()
    cols = list(items)
    df = pl.read_csv(d / "data-final.csv", separator="\t", columns=cols + ["IPC"], infer_schema_length=0)
    df = df.with_columns([pl.col(c).cast(pl.Int64, strict=False) for c in cols + ["IPC"]])
    df = df.filter(pl.col("IPC") == 1)
    df = df.filter(pl.all_horizontal([pl.col(c).is_between(1, 5) for c in cols]))
    n = df.height
    out = {}
    for c in cols:
        vc = dict(df.group_by(c).len().iter_rows())
        out[_norm(items[c])] = ({str(k - 1): vc.get(k, 0) / n for k in range(1, 6)}, n, c)
    return out


# ---------- normalize ----------

def normalize(raw_dir: Path) -> Iterator[Question]:
    parsers = [
        ("markers", _parse_markers),
        ("neo", _parse_neo),
        ("bfas", _parse_bfas),
        ("ab5c", _parse_ab5c),
        ("hexaco", _parse_hexaco),
    ]
    items: dict[str, dict] = {}
    order: list[str] = []
    for key, fn in parsers:
        for stem, trait, keyed, facet, scale in fn(raw_dir / PAGES[key]):
            text = _sentence(stem)
            k = _norm(text)
            if k in items:
                if scale not in items[k]["also_in"] and scale != items[k]["scale"]:
                    items[k]["also_in"].append(scale)
                continue
            items[k] = dict(text=text, trait=trait, keyed=keyed, facet=facet, scale=scale, group=key, also_in=[])
            order.append(k)

    # Markers, IPIP-NEO and BFAS are kept whole. Then fill to TARGET tier by tier (seeded):
    # AB5C, then HEXACO items on Big Five-like domains (X/A/C/O), then HEXACO H/E.
    tiers = [
        [k for k in order if items[k]["group"] in ("markers", "neo", "bfas")],
        [k for k in order if items[k]["group"] == "ab5c"],
        [k for k in order if items[k]["group"] == "hexaco" and items[k]["trait"]],
        [k for k in order if items[k]["group"] == "hexaco" and not items[k]["trait"]],
    ]
    rng = random.Random(SEED)
    chosen: list[str] = []
    for tier in tiers:
        room = TARGET - len(chosen)
        if room <= 0:
            break
        pick = set(tier) if len(tier) <= room else set(rng.sample(tier, room))
        chosen += [k for k in tier if k in pick]

    norms = _norms(raw_dir)
    for k in chosen:
        it = items[k]
        human = []
        meta = {"scale": it["scale"], "keyed": it["keyed"], "facet": it["facet"], "trait": it["trait"]}
        if it["also_in"]:
            meta["also_in"] = it["also_in"]
        if k in norms:
            dist, n, code = norms[k]
            human.append(
                HumanDist(
                    population="OpenPsychometrics web",
                    distribution=dist,
                    n=n,
                    source="Open Psychometrics IPIP-FFM-data-8Nov2018 (1=Disagree..5=Agree mapped to levels 0-4)",
                    wave="2016-2018",
                )
            )
            meta["openpsychometrics_item"] = code
        stmt = it["text"]
        yield Question(
            text=f'How well does this statement describe you: "{stmt}"',
            primitive="score",
            hemisphere="self",
            kind="personality",
            origin="dataset",
            source=NAME,
            options=LEVELS,
            node_hint=f"self.personality.big_five.{it['trait']}" if it["trait"] in TRAITS else "self.personality.big_five",
            human_text=f'How well would most people say this statement describes them: "{stmt}"',
            source_item_id=f"{it['scale']}|{it['facet'] or it['trait']}|{stmt}",
            license=LICENSE if not human else f"{LICENSE}; norms: {NORMS_LICENSE}",
            human=human,
            meta=meta,
        )
    # ---------- expansion (ASKJEV_TARGET_IPIP > 600): appended after the original sample ----------
    extra = TARGET_ALL - len(chosen)
    if extra > 0:
        yield from _extra(raw_dir, items, set(chosen), extra)


# ---------- expansion: the full IPIP pool ----------

BF = "self.personality.big_five."
P = "self.personality."
N_FEAR, N_TEMPER, N_SELF, N_IMP = BF + "neuroticism.fear_worry", BF + "neuroticism.temper_moodiness", \
    BF + "neuroticism.self_consciousness", BF + "neuroticism.impulses_coping"
E_SOC = BF + "extraversion.sociability_energy"
A_KIND, A_TRUST = BF + "agreeableness.kindness_cooperation", BF + "agreeableness.trust_modesty_temper"
C_DRIVE, C_ORDER = BF + "conscientiousness.drive_self_discipline", BF + "conscientiousness.order_caution"
O_IDEAS, O_CONV = BF + "openness.ideas_imagination", BF + "openness.convention_politics"
EMP, COPE, MIND = P + "emotions_stress.empathy_social_reading", P + "emotions_stress.coping_expressing", \
    P + "emotions_stress.mindfulness_awareness"

# Tedone construct label -> deepest fitting node. Unlisted labels fall back to the key-page trait, then
# self.personality.
LABEL_NODE = {
    # neuroticism
    "Anxiety": N_FEAR, "Fearfulness": N_FEAR, "Harm-Avoidance": N_FEAR, "Health Anxiety": N_FEAR, "Rumination": N_FEAR,
    "Vulnerability": N_FEAR, "Risk-avoidance": N_FEAR, "Calmness": BF + "neuroticism", "Tranquility": BF + "neuroticism",
    "Imperturbability": BF + "neuroticism", "Stability": BF + "neuroticism", "Emotional Stability": BF + "neuroticism",
    "Neuroticism": BF + "neuroticism", "Negative-Valence": BF + "neuroticism", "Security": BF + "neuroticism",
    "Cool-headedness": N_TEMPER,
    "Anger": N_TEMPER, "Affective Lability": N_TEMPER, "Emotionality": N_TEMPER, "Hostile Aggression": N_TEMPER,
    "Belligerence": N_TEMPER, "Toughness": BF + "neuroticism", "Sensitivity": N_TEMPER, "Patience": N_TEMPER,
    "Depression": BF + "neuroticism", "Anhedonia": BF + "neuroticism",
    "Self-consciousness": N_SELF, "Public Self-consciousness": N_SELF, "Conformity/Dependence/Need for approval": N_SELF,
    "Social-discomfort": N_SELF, "Timidity": N_SELF, "Relationship Insecurity": N_SELF, "Poise": N_SELF,
    "Immoderation": N_IMP, "Impulse-Control": N_IMP, "Moderation": N_IMP, "Temperance": N_IMP,
    "Self-control/Self-regulation": N_IMP, "Extravagance": N_IMP, "Recklessness": N_IMP, "Self-harm": N_IMP,
    # extraversion
    "Gregariousness": E_SOC, "Sociability": E_SOC, "Friendliness": E_SOC, "Warmth": E_SOC, "Cheerfulness": E_SOC,
    "Activity-Level": E_SOC, "Excitement-seeking": E_SOC, "Assertiveness": E_SOC, "Talkativeness": E_SOC,
    "Liveliness": E_SOC, "Social Boldness": E_SOC, "Social-confidence": E_SOC, "Vitality/Enthusiasm/Zest": E_SOC,
    "Joyfulness": E_SOC, "Positive Expressivity": E_SOC, "Expressiveness": E_SOC, "Forcefulness": E_SOC,
    "Dominance": E_SOC, "Leadership": E_SOC, "Extraversion": BF + "extraversion", "Introversion": BF + "extraversion",
    "Reserve": BF + "extraversion", "Reclusiveness": BF + "extraversion", "Social Withdrawal": BF + "extraversion",
    "Quickness": E_SOC, "Self-disclosure": BF + "extraversion", "Hypomanic Mood Intensity": E_SOC,
    # agreeableness
    "Altruism": A_KIND, "Sympathy": A_KIND, "Compassion": A_KIND, "Cooperation": A_KIND, "Tenderness": A_KIND,
    "Nurturance": A_KIND, "Kindness/Generosity": A_KIND, "Pleasantness": A_KIND, "Good Nature": A_KIND,
    "Amiability": A_KIND, "Gentleness": A_KIND, "Capacity for Love": A_KIND, "Teamwork/Citizenship": A_KIND,
    "Understanding": A_KIND, "Tolerance": A_KIND, "Flexibility": A_KIND, "Docility": A_KIND,
    "Forgiveness/Mercy": A_KIND, "Politeness": A_KIND,
    "Trust": A_TRUST, "Distrust": A_TRUST, "Mistrust": A_TRUST, "Modesty/Humility": A_TRUST, "Morality": A_TRUST,
    "Sincerity": A_TRUST, "Greed Avoidance": A_TRUST, "Equity/Fairness": A_TRUST, "Unpretentiousness": A_TRUST,
    "Honesty/Integrity/Authenticity": A_TRUST, "Provocativeness": A_TRUST, "Rudeness": A_TRUST,
    "Agreeableness": BF + "agreeableness",
    # conscientiousness
    "Achievement-striving": C_DRIVE, "Self-discipline": C_DRIVE, "Industriousness/Perseverance/Persistence": C_DRIVE,
    "Diligence": C_DRIVE, "Purposefulness": C_DRIVE, "Self-efficacy": C_DRIVE, "Competence": C_DRIVE,
    "Efficiency": C_DRIVE, "Initiative": C_DRIVE, "Dutifulness": C_DRIVE, "Responsibility": C_DRIVE,
    "Non-Perseverance": C_DRIVE, "Irresponsibility": C_DRIVE, "Resourcefulness": C_DRIVE,
    "Orderliness": C_ORDER, "Organization": C_ORDER, "Methodicalness": C_ORDER, "Cautiousness": C_ORDER,
    "Deliberateness": C_ORDER, "Prudence": C_ORDER, "Planfulness": C_ORDER, "Non-Planfulness": C_ORDER,
    "Perfectionism": C_ORDER, "Need for Order and Cleanliness": C_ORDER, "Disorderliness": C_ORDER,
    "Reflection": C_ORDER, "Rationality": C_ORDER, "Conscientiousness": BF + "conscientiousness",
    # openness
    "Imagination": O_IDEAS, "Intellect": O_IDEAS, "Ingenuity": O_IDEAS, "Creativity/Originality": O_IDEAS,
    "Curiosity": O_IDEAS, "Inquisitiveness": O_IDEAS, "Intellectual Openness": O_IDEAS, "Intellectual-Breadth": O_IDEAS,
    "Intellectual-Complexity": O_IDEAS, "Complexity": O_IDEAS, "Depth": O_IDEAS, "Love of Learning": O_IDEAS,
    "Love of Reading": O_IDEAS, "Need for Cognition": O_IDEAS, "Language Mastery": O_IDEAS, "Comprehension": O_IDEAS,
    "Problem-solving": O_IDEAS, "Fantasy Proneness": O_IDEAS, "Culture": O_IDEAS, "Science Interest": O_IDEAS,
    "Liberalism": O_CONV, "Conservatism": O_CONV, "Traditionalism": O_CONV, "Unconventionality": O_CONV,
    "Rebelliousness": O_CONV, "Variety-seeking": BF + "openness", "Adventurousness": BF + "openness",
    "Aesthetic Appreciation/Artistic Interests": BF + "openness", "Openness To Experience": BF + "openness",
    "Judgment/Open-mindedness": BF + "openness",
    # emotions
    "Empathy": EMP, "Social/Personal/Emotional Intelligence": EMP, "Responsive Distress": EMP, "Responsive Joy": EMP,
    "Insight": EMP, "Attention to Emotions": MIND, "Introspection/Private Self-Consciousness": MIND,
    "Emotion-based Decision-making": COPE, "Negative Expressivity": COPE, "Sentimentality": COPE,
    "Emotional Detachment": COPE, "Romantic Disinterest": COPE,
    # other personality areas
    "Grandiosity": P + "dark_side", "Manipulativeness": P + "dark_side", "Machiavellianism": P + "dark_side",
    "Callousness": P + "dark_side", "Exhibitionism": P + "dark_side", "Disparagement": P + "dark_side",
    "Norm Violation": P + "dark_side", "Hypomanic Exhibitionism": P + "dark_side", "Submissiveness": P + "dark_side",
    "Humor/Playfulness": P + "humor_style.humor_habits_tastes",
    "Ambition/Drive": P + "motivation_ambition", "Workaholism": P + "motivation_ambition",
    "Behavioral Inhibition/Activation System": P + "motivation_ambition", "Independence": P + "motivation_ambition",
    "Self-Sufficiency": P + "motivation_ambition", "Adaptability": P + "motivation_ambition",
    "Risk-taking/Sensation-Seeking/Thrill-Seeking": P + "risk_decision_style",
    "Locus of Control": P + "risk_decision_style", "Locus of Control,Chance": P + "risk_decision_style",
    "Locus of Control,External": P + "risk_decision_style", "Locus of Control,Internal": P + "risk_decision_style",
    "Locus of Control,Rational": P + "risk_decision_style",
    "Self-esteem": P + "self_concept", "Self-acceptance": P + "self_concept", "Self-confidence": P + "self_concept",
    "Attractiveness": P + "self_concept", "Physical Attractiveness": P + "self_concept",
    "Appearance-Consciousness": P + "self_concept", "Femininity": P + "self_concept", "Self-monitoring": P + "self_concept",
    "Impression-Management": P + "self_concept", "Self-deception": P + "self_concept", "Unlikely Virtues": P + "self_concept",
    "Peculiarity": P + "self_concept", "Rigidity": C_ORDER, "Cognitive-Failures": C_ORDER, "ADHD": C_DRIVE,
    "Happiness": "self.mind.happiness_wellbeing", "Satisfaction": "self.mind.happiness_wellbeing",
    "Hope/Optimism": "self.mind.happiness_wellbeing", "Gratitude": "self.mind.happiness_wellbeing",
    "Wisdom/Perspective": "self.mind.big_questions.meaning_human_nature",
    "Spirituality/Religiousness": "self.mind.big_questions", "Irrational Beliefs": "self.mind.luck_fate",
    "Bravery/Courage/Valor": P + "risk_decision_style", "Romanticism": "self.love.romance_partnership",
    "Interest in Romance": "self.love.romance_partnership", "Interest in Children": "self.love.family_parenting",
    "Interest in Environmentalism": "self.values.animals_environment", "Interest in Pets": "self.values.animals_environment",
    "Interest in Political Activism": O_CONV, "Interest in Money": P + "motivation_ambition",
    "Interest in Self-Improvement": P + "motivation_ambition",
    **{f"Interest in {x}": P + "interests" for x in (
        "Collecting", "Computing", "Drinking", "Exercise", "Food", "Gambling", "Game-Playing", "Gardening", "Home-Making",
        "Journaling", "Music", "Outdoor Activities", "Shopping", "Social Media", "Sports", "Travel", "Vehicles",
        "Watching Television")},
    "Construction/Mechanical Interests": P + "interests",
}
# Labels / instruments of clinical symptom scales (CES-D, dissociation, OCD symptoms, CAT-PD pathology scales):
# kept but flagged sensitive.
CLINICAL_LABELS = {"Self-harm", "Dissociation", "Unusual Experiences", "Obsessive-Compulsive Symptoms", "Anhedonia",
                   "Cognitive Problems", "Health Anxiety"}
CLINICAL_INSTRUMENTS = {"Radloff1977", "Foa2002"}
# Label priority when an item is scored on several scales: facet-level inventories first.
EMPIRICAL = {"CPI", "MPQ", "JPI", "BIS_BAS", "7FACTOR"}
INSTRUMENT_ORDER = ["NEO", "BFAS", "AB5C", "HEXACO_PI", "16PF", "6FPQ", "CPI", "HPI-HIC", "HPI", "JPI", "MPQ", "TCI",
                    "VIA", "ORAIS", "ORVIS", "IPIP-IPC", "BIS_BAS", "Barchard2001", "CAT-PD"]
SENSITIVE = re.compile(r"(?i)\b(sex\w*|suicid\w*|kill(ing)? myself|hurt(ing)? myself|self-inflicted|cut myself|"
                       r"harm myself|voices inside my head|thoughts about death|drugs?|porn\w*|naked|nude)\b")
POLITICAL = re.compile(r"(?i)\b(vote|voting|liberal|conservative|politic\w*|religio\w*|god|church|pray\w*|faith|scriptures?|"
                       r"abortion|immigra\w*|races and religions|death penalty|national anthem|patriot\w*)\b")
CLAUSE = re.compile(r"^(When|As|In|If|After|Before|Whenever|Once)\b[^,]*, ")
CODE = re.compile(r"[A-Z]\d+ ?\*?(, [A-Z]\d+ ?\*?)*")


def _alpha_items(path: Path) -> list[tuple[str, str]]:
    """(stem, survey codes) pairs from the alphabetical item list: each item line is followed by its codes."""
    ls = _lines(path)
    out = []
    for i in range(len(ls) - 1):
        if CODE.fullmatch(ls[i + 1]) and not CODE.fullmatch(ls[i]) and len(ls[i]) < 200 and ls[i][:1].isupper():
            out.append((ls[i], ls[i + 1]))
    return out


def _tedone(path: Path) -> dict[str, list[tuple[str, str, int | None]]]:
    import openpyxl  # only needed for the expansion

    wb = openpyxl.load_workbook(path, read_only=True)
    rows = list(wb.worksheets[0].iter_rows(values_only=True))
    hdr = [str(h) for h in rows[0]]
    ii, ki, ti, li = hdr.index("instrument"), hdr.index("key"), hdr.index("text"), hdr.index("label")
    out: dict[str, list] = {}
    for r in rows[1:]:
        if not r[ti] or not r[li] or r[li] == "None":
            continue
        out.setdefault(_norm(_extra_sentence(str(r[ti]))), []).append((str(r[ii]), str(r[li]), r[ki]))
    return out


def _extra_sentence(stem: str) -> str:
    """_sentence, plus stems that open with a clause ("When in a group, try ...") get "I" after the clause."""
    s = re.sub(r"\s+", " ", stem.strip())
    m = CLAUSE.match(s)
    if m:
        body = _sentence(s[m.end():])
        return s[: m.end()] + body
    return _sentence(s)


def _pick_label(labels: list[tuple[str, str, int | None]]) -> tuple[str, str, int | None] | None:
    """The label that sets node_hint: first mapped label from a content-homogeneous inventory. Labels of the
    IPIP scales built to correlate with CPI, MPQ, JPI, BIS/BAS or the 7-factor model are kept in meta but not
    used for placement: their items are often off-topic for the label (CPI "Adventurousness" has
    "Believe that I am important.")."""
    ranked = sorted(labels, key=lambda x: INSTRUMENT_ORDER.index(x[0]) if x[0] in INSTRUMENT_ORDER else 99)
    for lab in ranked:
        if lab[1] in LABEL_NODE and lab[0] not in EMPIRICAL:
            return lab
    return None


def _loose(k: str) -> str:
    return re.sub(r"\b(that|am|i)\b| ", "", k)


def _match_labels(ted: dict[str, list], keys: list[str]) -> dict[str, list]:
    """Tedone's texts sometimes drop "that" or "Am" or change a word; match exact, then loosely, then fuzzily."""
    import difflib

    loose = {}
    for k in ted:
        loose.setdefault(_loose(k), []).extend(ted[k])
    out = {}
    lk = list(loose)
    for k in keys:
        if k in ted:
            out[k] = ted[k]
        elif _loose(k) in loose:
            out[k] = loose[_loose(k)]
        else:
            m = difflib.get_close_matches(_loose(k), lk, 1, 0.92)
            if m:
                out[k] = loose[m[0]]
    return out


def _extra(raw_dir: Path, items: dict[str, dict], taken: set[str], k: int) -> Iterator[Question]:
    ted = _tedone(raw_dir / TEDONE_XLSX)
    pool: dict[str, dict] = {}
    # 1) key-page items outside the original sample (AB5C / HEXACO leftovers), with their parsed trait/facet
    for key, it in items.items():
        if key not in taken:
            pool[key] = dict(it, code=None)
    # 2) every other item on the alphabetical list
    for stem, code in _alpha_items(raw_dir / ALPHA_PAGE):
        text = _extra_sentence(stem)
        key = _norm(text)
        if key in taken:
            continue
        if key in pool:
            pool[key]["code"] = pool[key]["code"] or code
            continue
        pool[key] = dict(text=text, trait=None, keyed=None, facet=None, scale=None, group="alpha", also_in=[], code=code)

    picked = hash_order(list(pool), key=lambda x: x, salt=EXTRA_SALT)[:k]
    matched = _match_labels(ted, picked)
    for key in picked:
        it = pool[key]
        labels = matched.get(key, [])
        lab = _pick_label(labels)
        node = LABEL_NODE.get(lab[1]) if lab else None
        if node is None:
            node = f"{BF}{it['trait']}" if it["trait"] in TRAITS else ("self.personality.big_five" if it["scale"] else P[:-1])
        trait = next((t for t in TRAITS if node.startswith(BF + t)), it["trait"])
        if it["scale"]:  # parsed from a scoring-key page: keep that scale/facet/key
            scale, facet, keyed = it["scale"], it["facet"], it["keyed"]
        elif labels:
            lab0 = lab or labels[0]
            scale, facet = f"{lab0[0]} (IPIP)", lab0[1]
            keyed = {1: "+", -1: "-"}.get(lab0[2])
        else:
            scale, facet, keyed = "IPIP item pool (not scored on a published IPIP scale)", None, None
        stmt = it["text"]
        flags = []
        if SENSITIVE.search(stmt) or any(l in CLINICAL_LABELS or i in CLINICAL_INSTRUMENTS for i, l, _ in labels):
            flags.append("sensitive")
        if POLITICAL.search(stmt) or any(l in ("Liberalism", "Interest in Political Activism")
                                         for _, l, _ in labels):
            flags.append("political")
        meta = {"scale": scale, "keyed": keyed, "facet": facet, "trait": trait}
        if it["also_in"]:
            meta["also_in"] = it["also_in"]
        if it.get("code"):
            meta["ipip_survey_codes"] = it["code"]
        if labels:
            meta["ipip_scales"] = [{"instrument": i, "label": l, "keyed": {1: "+", -1: "-"}.get(kk)} for i, l, kk in labels]
        if flags:
            meta["flags"] = flags
        yield Question(
            text=f'How well does this statement describe you: "{stmt}"',
            primitive="score",
            hemisphere="self",
            kind="personality",
            origin="dataset",
            source=NAME,
            options=LEVELS,
            node_hint=node,
            human_text=f'How well would most people say this statement describes them: "{stmt}"',
            source_item_id=f"{scale}|{facet or trait or 'unscored'}|{stmt}",
            license=LICENSE,
            meta=meta,
        )
